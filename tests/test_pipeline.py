"""Reader, extractor (fake LLM), scoring and export."""

import io
import json

import pytest
from openpyxl import load_workbook

from app.evaluation import field_correct
from app.export import ApprovedInvoice, to_csv, to_excel
from app.extract import choose_mode, extract
from app.reader import UnsupportedFile, read_document, read_path
from tests.conftest import INVOICES, FakeLLM, make_invoice

# ----------------------------------------------------------------------- reader

def test_digital_pdf_has_text_and_pages():
    doc = read_path(INVOICES / "inv01-us-pdf.pdf")
    assert doc.kind == "pdf_text" and "INVOICE" in doc.text and len(doc.pages) == 1


def test_scanned_pdf_has_no_text():
    doc = read_path(INVOICES / "inv21-us-scan.pdf")
    assert doc.kind == "pdf_scan" and doc.text is None and doc.pages


def test_multipage_pdf_marks_pages():
    doc = read_path(INVOICES / "inv10-de-pdf.pdf")
    assert len(doc.pages) == 2 and "--- page 2 ---" in doc.text


def test_photo_is_resized_jpeg():
    doc = read_path(INVOICES / "inv27-uk-photo.jpg")
    assert doc.kind == "image" and doc.pages[0].mime_type == "image/jpeg"


@pytest.mark.parametrize(("data", "name"), [(b"hello", "a.txt"), (b"not a pdf", "a.pdf"), (b"nope", "a.png")])
def test_unsupported_files(data, name):
    with pytest.raises(UnsupportedFile):
        read_document(data, name)


# ----------------------------------------------------------------------- extract

def test_auto_mode_choice():
    assert choose_mode(read_path(INVOICES / "inv01-us-pdf.pdf"), "auto") == "hybrid"
    assert choose_mode(read_path(INVOICES / "inv21-us-scan.pdf"), "auto") == "vision"
    assert choose_mode(read_path(INVOICES / "inv21-us-scan.pdf"), "text") == "vision"


def test_extract_sends_text_and_images_in_hybrid(invoice):
    llm = FakeLLM(invoice.model_dump_json())
    ex = extract(read_path(INVOICES / "inv01-us-pdf.pdf"), "auto", llm)
    assert ex.invoice == invoice and ex.mode == "hybrid"
    assert llm.calls[0]["images"] == 1 and "Invoice #" in llm.calls[0]["prompt"]


def test_extract_retries_once_on_schema_error(invoice):
    llm = FakeLLM('{"vendor_name": "x"}', invoice.model_dump_json())
    ex = extract(read_path(INVOICES / "inv01-us-pdf.pdf"), "text", llm)
    assert ex.invoice == invoice and ex.attempts == 2
    assert "did not match the schema" in llm.calls[1]["prompt"]


def test_extract_gives_up_after_two_bad_answers():
    ex = extract(read_path(INVOICES / "inv01-us-pdf.pdf"), "text", FakeLLM("{}", "not json"))
    assert ex.invoice is None and ex.error


# ----------------------------------------------------------------------- scoring

def test_field_correct_rules():
    assert field_correct("vendor_name", "KESTREL PRINT STUDIO LTD.", "Kestrel Print Studio Ltd")
    assert field_correct("total", 1080.004, 1080.0)
    assert not field_correct("total", 1080.5, 1080.0)
    assert field_correct("invoice_number", "AL 7781", "AL-7781") is False
    assert field_correct("due_date", None, None)
    assert field_correct("line_items", [{"amount": 2}, {"amount": 1}], [{"amount": 1}, {"amount": 2.0}])


def test_truth_is_consistent(truth):
    assert len(truth) == 30
    for t in truth.values():
        assert abs(t["subtotal"] + t["tax"] - t["total"]) < 0.011
        assert t["amount_due"] <= t["total"]
        assert (INVOICES / t["file"]).exists()


# ----------------------------------------------------------------------- export

def test_excel_export_has_typed_cells():
    approved = [ApprovedInvoice("a.pdf", make_invoice(), corrected_fields=["total"])]
    wb = load_workbook(io.BytesIO(to_excel(approved)))
    ws = wb["Invoices"]
    assert ws["B2"].value == "Northwind Office Supply"
    assert ws["J2"].value == 1353.13 and ws["J2"].number_format == "#,##0.00"
    assert ws["E2"].value.year == 2026
    assert wb["Line items"].max_row == 3


def test_csv_export():
    text = to_csv([ApprovedInvoice("a.pdf", make_invoice())])
    assert "2026-03-04" in text and "INV-1042" in text and json.dumps("a.pdf")[1:-1] in text
